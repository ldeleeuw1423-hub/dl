"""Export service — generates Excel and PDF reports for project estimations.

Uses openpyxl for Excel and reportlab for PDF.
"""
import io
import logging
from datetime import date, datetime
from typing import Optional

from sqlalchemy.orm import Session

from app.models.estimation import Estimation
from app.models.permit import Permit
from app.models.project import Project
from app.models.risk import Risk

logger = logging.getLogger(__name__)

# Colour palette
NAVY = "1E3A5F"
NAVY_LIGHT = "2C4D7A"
GREY_ROW = "F2F5F9"
WHITE = "FFFFFF"
GREEN = "27AE60"
AMBER = "F39C12"
RED_C = "E74C3C"

RISK_LEVEL_COLOUR = {
    "low": "27AE60",
    "medium": "F39C12",
    "high": "E74C3C",
    "critical": "8E1A1A",
}

STATUS_NL = {
    "required": "Vereist",
    "in_preparation": "In voorbereiding",
    "submitted": "Ingediend",
    "approved": "Goedgekeurd",
    "rejected": "Afgewezen",
    "not_required": "Niet vereist",
    "open": "Open",
    "mitigated": "Gemitigeerd",
    "accepted": "Geaccepteerd",
    "closed": "Gesloten",
    "active": "Actief",
    "on_hold": "On hold",
    "completed": "Afgerond",
    "cancelled": "Geannuleerd",
}

DISCIPLINE_NL = {
    "Gas": "Gas",
    "Elektra": "Elektra",
    "LS_MS": "LS/MS",
    "Stations": "Stations",
}

PHASE_NL = {
    "VO": "Voorontwerp",
    "DO": "Definitief Ontwerp",
    "UO": "Uitvoeringsontwerp",
    "Realisatie": "Realisatie",
}


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _fmt_currency(value) -> str:
    if value is None:
        return "—"
    return f"€ {float(value):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def _fmt_number(value, decimals: int = 1) -> str:
    if value is None:
        return "—"
    fmt = f"{{:.{decimals}f}}"
    return fmt.format(float(value))


def _fmt_date(value) -> str:
    if value is None:
        return "—"
    if isinstance(value, (date, datetime)):
        return value.strftime("%d-%m-%Y")
    return str(value)


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------


def generate_estimation_excel(project_id: str, db: Session) -> bytes:
    """Generate an Excel workbook for a project and return raw bytes."""
    from openpyxl import Workbook
    from openpyxl.styles import (
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )
    from openpyxl.utils import get_column_letter

    project: Optional[Project] = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise ValueError(f"Project {project_id} not found")

    estimation: Optional[Estimation] = (
        db.query(Estimation)
        .filter(Estimation.project_id == project.id, Estimation.is_current == True)  # noqa: E712
        .order_by(Estimation.version.desc())
        .first()
    )

    risks = (
        db.query(Risk)
        .filter(Risk.project_id == project.id)
        .order_by(Risk.score.desc())
        .all()
    )

    permits = (
        db.query(Permit)
        .filter(Permit.project_id == project.id)
        .order_by(Permit.risk_level.desc())
        .all()
    )

    wb = Workbook()
    wb.remove(wb.active)  # remove default sheet

    # -- Styles --
    header_fill = PatternFill("solid", fgColor=NAVY)
    subheader_fill = PatternFill("solid", fgColor=NAVY_LIGHT)
    alt_fill = PatternFill("solid", fgColor=GREY_ROW)
    white_fill = PatternFill("solid", fgColor=WHITE)
    header_font = Font(bold=True, color=WHITE, size=10)
    bold_font = Font(bold=True, size=10)
    normal_font = Font(size=10)
    thin_border = Border(
        left=Side(style="thin", color="D0D8E4"),
        right=Side(style="thin", color="D0D8E4"),
        top=Side(style="thin", color="D0D8E4"),
        bottom=Side(style="thin", color="D0D8E4"),
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    def _style_header_row(ws, row: int, num_cols: int) -> None:
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = thin_border

    def _style_data_row(ws, row: int, num_cols: int, alternate: bool) -> None:
        fill = alt_fill if alternate else white_fill
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = fill
            cell.font = normal_font
            cell.alignment = left_align
            cell.border = thin_border

    def _auto_width(ws) -> None:
        for col_cells in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col_cells[0].column)
            for cell in col_cells:
                try:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 60)

    # -------------------------------------------------------------------------
    # Sheet 1 — Urenschatter
    # -------------------------------------------------------------------------
    ws1 = wb.create_sheet("Urenschatter")
    ws1.sheet_view.showGridLines = False
    ws1.row_dimensions[1].height = 20

    headers1 = [
        "Discipline",
        "Uren Engineering",
        "Uren PM",
        "Uren OM",
        "Uren Werkvoorbereiding",
        "Uren Uitvoering",
        "Totaal Uren",
        "Bandbreedte Laag",
        "Bandbreedte Hoog",
        "Betrouwbaarheidsscore",
        "Methodiek",
        "Onderbouwing",
    ]
    for col_idx, h in enumerate(headers1, 1):
        ws1.cell(row=1, column=col_idx, value=h)
    _style_header_row(ws1, 1, len(headers1))

    if estimation:
        total_hours = sum(
            float(v or 0)
            for v in [
                estimation.hours_engineering,
                estimation.hours_pm,
                estimation.hours_om,
                estimation.hours_workprep,
                estimation.hours_execution,
            ]
        )
        row_data = [
            DISCIPLINE_NL.get(estimation.discipline or project.discipline, estimation.discipline or project.discipline),
            _fmt_number(estimation.hours_engineering),
            _fmt_number(estimation.hours_pm),
            _fmt_number(estimation.hours_om),
            _fmt_number(estimation.hours_workprep),
            _fmt_number(estimation.hours_execution),
            _fmt_number(total_hours),
            _fmt_currency(estimation.bandwidth_low),
            _fmt_currency(estimation.bandwidth_high),
            f"{_fmt_number(estimation.confidence_score, 0)}%",
            estimation.methodology or "—",
            estimation.reasoning or "—",
        ]
        for col_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=2, column=col_idx, value=val)
            cell.fill = alt_fill
            cell.font = normal_font
            cell.alignment = left_align
            cell.border = thin_border

        # Totals row
        total_row = 3
        ws1.cell(row=total_row, column=1, value="TOTAAL")
        ws1.cell(row=total_row, column=7, value=_fmt_number(total_hours))
        ws1.cell(row=total_row, column=8, value=_fmt_currency(estimation.bandwidth_low))
        ws1.cell(row=total_row, column=9, value=_fmt_currency(estimation.bandwidth_high))
        for col in range(1, len(headers1) + 1):
            cell = ws1.cell(row=total_row, column=col)
            cell.fill = subheader_fill
            cell.font = Font(bold=True, color=WHITE, size=10)
            cell.border = thin_border
            cell.alignment = center_align
    else:
        ws1.cell(row=2, column=1, value="Geen raming beschikbaar")
        ws1.cell(row=2, column=1).font = Font(italic=True, color="888888")

    _auto_width(ws1)

    # -------------------------------------------------------------------------
    # Sheet 2 — Risicologboek
    # -------------------------------------------------------------------------
    ws2 = wb.create_sheet("Risicologboek")
    ws2.sheet_view.showGridLines = False

    headers2 = [
        "Nr.",
        "Beschrijving",
        "Categorie",
        "Kans (1-5)",
        "Impact (1-5)",
        "Score",
        "Maatregel",
        "Eigenaar",
        "Deadline",
        "Status",
        "Restrisico",
        "Bron",
    ]
    for col_idx, h in enumerate(headers2, 1):
        ws2.cell(row=1, column=col_idx, value=h)
    _style_header_row(ws2, 1, len(headers2))

    for r_idx, risk in enumerate(risks, 1):
        row = r_idx + 1
        _style_data_row(ws2, row, len(headers2), r_idx % 2 == 0)
        values = [
            r_idx,
            risk.description,
            risk.category,
            risk.probability,
            risk.impact,
            risk.score,
            risk.mitigation_measure or "—",
            risk.owner or "—",
            _fmt_date(risk.deadline),
            STATUS_NL.get(risk.status, risk.status),
            risk.residual_risk if risk.residual_risk is not None else "—",
            risk.source,
        ]
        for col_idx, val in enumerate(values, 1):
            ws2.cell(row=row, column=col_idx).value = val

        # Colour score cell based on value
        score_cell = ws2.cell(row=row, column=6)
        if risk.score >= 15:
            score_cell.fill = PatternFill("solid", fgColor="E74C3C")
            score_cell.font = Font(bold=True, color=WHITE, size=10)
        elif risk.score >= 9:
            score_cell.fill = PatternFill("solid", fgColor="F39C12")
            score_cell.font = Font(bold=True, size=10)
        else:
            score_cell.fill = PatternFill("solid", fgColor="27AE60")
            score_cell.font = Font(bold=True, color=WHITE, size=10)

    if not risks:
        ws2.cell(row=2, column=1, value="Geen risico's geregistreerd").font = Font(italic=True, color="888888")

    _auto_width(ws2)

    # -------------------------------------------------------------------------
    # Sheet 3 — Vergunningen
    # -------------------------------------------------------------------------
    ws3 = wb.create_sheet("Vergunningen")
    ws3.sheet_view.showGridLines = False

    headers3 = [
        "Nr.",
        "Type vergunning",
        "Omschrijving",
        "Bevoegd gezag",
        "Status",
        "Indiendatum",
        "Verwachte goedkeuring",
        "Daadwerkelijke goedkeuring",
        "Risiconiveau",
        "Vertragingskans (%)",
        "Eigenaar",
        "Opmerkingen",
    ]
    for col_idx, h in enumerate(headers3, 1):
        ws3.cell(row=1, column=col_idx, value=h)
    _style_header_row(ws3, 1, len(headers3))

    for p_idx, permit in enumerate(permits, 1):
        row = p_idx + 1
        _style_data_row(ws3, row, len(headers3), p_idx % 2 == 0)
        values = [
            p_idx,
            permit.permit_type,
            permit.description or "—",
            permit.authority,
            STATUS_NL.get(permit.status, permit.status),
            _fmt_date(permit.submission_date),
            _fmt_date(permit.expected_approval),
            _fmt_date(permit.actual_approval),
            permit.risk_level.upper() if permit.risk_level else "—",
            permit.delay_probability if permit.delay_probability is not None else "—",
            permit.owner or "—",
            permit.notes or "—",
        ]
        for col_idx, val in enumerate(values, 1):
            ws3.cell(row=row, column=col_idx).value = val

        # Colour risk_level cell
        rl_cell = ws3.cell(row=row, column=9)
        colour = RISK_LEVEL_COLOUR.get(permit.risk_level, WHITE)
        rl_cell.fill = PatternFill("solid", fgColor=colour)
        rl_cell.font = Font(bold=True, color=WHITE, size=10)

    if not permits:
        ws3.cell(row=2, column=1, value="Geen vergunningen geregistreerd").font = Font(italic=True, color="888888")

    _auto_width(ws3)

    # -------------------------------------------------------------------------
    # Sheet 4 — Samenvatting
    # -------------------------------------------------------------------------
    ws4 = wb.create_sheet("Samenvatting")
    ws4.sheet_view.showGridLines = False

    def _kv(ws, row: int, key: str, value: str) -> None:
        kc = ws.cell(row=row, column=1, value=key)
        kc.font = Font(bold=True, size=10)
        kc.fill = alt_fill
        kc.border = thin_border
        kc.alignment = left_align
        vc = ws.cell(row=row, column=2, value=value)
        vc.font = normal_font
        vc.border = thin_border
        vc.alignment = left_align

    # Title
    title_cell = ws4.cell(row=1, column=1, value="PROJECTSAMENVATTING")
    title_cell.font = Font(bold=True, color=WHITE, size=13)
    title_cell.fill = header_fill
    title_cell.alignment = center_align
    ws4.merge_cells("A1:B1")
    ws4.row_dimensions[1].height = 25

    _kv(ws4, 2, "Projectnummer", project.project_number)
    _kv(ws4, 3, "Projectnaam", project.name)
    _kv(ws4, 4, "Opdrachtgever", project.client or "—")
    _kv(ws4, 5, "Locatie", project.location or "—")
    _kv(ws4, 6, "Discipline", DISCIPLINE_NL.get(project.discipline, project.discipline))
    _kv(ws4, 7, "Fase", PHASE_NL.get(project.phase, project.phase))
    _kv(ws4, 8, "Status", STATUS_NL.get(project.status, project.status))
    _kv(ws4, 9, "Startdatum", _fmt_date(project.start_date))
    _kv(ws4, 10, "Einddatum", _fmt_date(project.end_date))
    _kv(ws4, 11, "Budget geraamd", _fmt_currency(project.budget_estimated))
    _kv(ws4, 12, "Budget actueel", _fmt_currency(project.budget_actual))
    _kv(ws4, 13, "Uren geraamd", f"{_fmt_number(project.hours_estimated, 0)} uur" if project.hours_estimated else "—")
    _kv(ws4, 14, "Uren actueel", f"{_fmt_number(project.hours_actual, 0)} uur" if project.hours_actual else "—")
    _kv(ws4, 15, "Aantal risico's", str(len(risks)))
    _kv(ws4, 16, "Hoog/kritiek risico's", str(sum(1 for r in risks if r.score >= 15)))
    _kv(ws4, 17, "Aantal vergunningen", str(len(permits)))
    _kv(ws4, 18, "Vergunningen hoog risico", str(sum(1 for p in permits if p.risk_level in ("high", "critical"))))
    _kv(ws4, 19, "Rapportdatum", datetime.now().strftime("%d-%m-%Y %H:%M"))

    ws4.column_dimensions["A"].width = 32
    ws4.column_dimensions["B"].width = 45

    # -------------------------------------------------------------------------
    # Serialize to bytes
    # -------------------------------------------------------------------------
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.read()


# ---------------------------------------------------------------------------
# PDF export
# ---------------------------------------------------------------------------


def generate_estimation_pdf(project_id: str, db: Session) -> bytes:
    """Generate a PDF report for a project and return raw bytes."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    project: Optional[Project] = db.query(Project).filter(Project.id == project_id).first()
    if project is None:
        raise ValueError(f"Project {project_id} not found")

    estimation: Optional[Estimation] = (
        db.query(Estimation)
        .filter(Estimation.project_id == project.id, Estimation.is_current == True)  # noqa: E712
        .order_by(Estimation.version.desc())
        .first()
    )

    risks = (
        db.query(Risk)
        .filter(Risk.project_id == project.id)
        .order_by(Risk.score.desc())
        .limit(10)
        .all()
    )

    permits = (
        db.query(Permit)
        .filter(Permit.project_id == project.id)
        .order_by(Permit.risk_level.desc())
        .all()
    )

    buffer = io.BytesIO()

    navy_color = colors.HexColor(f"#{NAVY}")
    light_grey = colors.HexColor("#F2F5F9")
    mid_grey = colors.HexColor("#9BA8B8")
    white_c = colors.white

    page_width, page_height = A4
    margin_left = 20 * mm
    margin_right = 20 * mm
    margin_top = 25 * mm
    margin_bottom = 20 * mm

    report_date = datetime.now().strftime("%d-%m-%Y")

    def _on_page(canvas, doc):
        """Header & footer on every page."""
        canvas.saveState()
        # Header bar
        canvas.setFillColor(navy_color)
        canvas.rect(0, page_height - 15 * mm, page_width, 15 * mm, fill=1, stroke=0)
        canvas.setFillColor(white_c)
        canvas.setFont("Helvetica-Bold", 9)
        canvas.drawString(margin_left, page_height - 9 * mm, "PROJECTRAPPORTAGE — ONDERGRONDSE INFRASTRUCTUUR")
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(
            page_width - margin_right,
            page_height - 9 * mm,
            f"{project.project_number}  |  {project.name}",
        )

        # Footer
        canvas.setFillColor(mid_grey)
        canvas.setFont("Helvetica", 7)
        canvas.drawString(margin_left, 10 * mm, f"Gegenereerd op {report_date}")
        canvas.drawRightString(
            page_width - margin_right,
            10 * mm,
            f"Pagina {doc.page}",
        )
        canvas.setStrokeColor(navy_color)
        canvas.setLineWidth(0.5)
        canvas.line(margin_left, 14 * mm, page_width - margin_right, 14 * mm)
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=margin_left,
        rightMargin=margin_right,
        topMargin=margin_top,
        bottomMargin=margin_bottom,
        onFirstPage=_on_page,
        onLaterPages=_on_page,
    )

    styles = getSampleStyleSheet()
    style_h1 = ParagraphStyle(
        "H1",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=16,
        textColor=navy_color,
        spaceAfter=4 * mm,
    )
    style_h2 = ParagraphStyle(
        "H2",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=12,
        textColor=navy_color,
        spaceBefore=6 * mm,
        spaceAfter=3 * mm,
    )
    style_body = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=13,
        spaceAfter=2 * mm,
    )
    style_label = ParagraphStyle(
        "Label",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=9,
        textColor=navy_color,
    )
    style_small = ParagraphStyle(
        "Small",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        textColor=colors.HexColor("#555555"),
    )

    table_header_style = [
        ("BACKGROUND", (0, 0), (-1, 0), navy_color),
        ("TEXTCOLOR", (0, 0), (-1, 0), white_c),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, 0), 8),
        ("ALIGN", (0, 0), (-1, -1), "LEFT"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [white_c, light_grey]),
        ("FONTSIZE", (0, 1), (-1, -1), 8),
        ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D0D8E4")),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]

    story = []

    # -------------------------------------------------------------------------
    # Cover / Title section
    # -------------------------------------------------------------------------
    story.append(Paragraph(project.name, style_h1))
    story.append(
        Paragraph(
            f"Projectnummer: <b>{project.project_number}</b> &nbsp;&nbsp; "
            f"Discipline: <b>{DISCIPLINE_NL.get(project.discipline, project.discipline)}</b> &nbsp;&nbsp; "
            f"Fase: <b>{PHASE_NL.get(project.phase, project.phase)}</b>",
            style_body,
        )
    )
    if project.client:
        story.append(Paragraph(f"Opdrachtgever: {project.client}", style_body))
    if project.location:
        story.append(Paragraph(f"Locatie: {project.location}", style_body))
    story.append(HRFlowable(width="100%", thickness=1, color=navy_color, spaceAfter=4 * mm))

    # -------------------------------------------------------------------------
    # Executive summary
    # -------------------------------------------------------------------------
    story.append(Paragraph("Managementsamenvatting", style_h2))

    summary_data = [
        ["KPI", "Waarde"],
        ["Budget geraamd", _fmt_currency(project.budget_estimated)],
        ["Budget actueel", _fmt_currency(project.budget_actual)],
        ["Uren geraamd", f"{_fmt_number(project.hours_estimated, 0)} uur"],
        ["Uren actueel", f"{_fmt_number(project.hours_actual, 0)} uur"],
        ["Startdatum", _fmt_date(project.start_date)],
        ["Einddatum", _fmt_date(project.end_date)],
        ["Status", STATUS_NL.get(project.status, project.status)],
        ["Totaal risico's", str(len(risks))],
        ["Kritieke risico's (score ≥ 15)", str(sum(1 for r in risks if r.score >= 15))],
        ["Totaal vergunningen", str(len(permits))],
    ]

    if estimation:
        summary_data.append(["Betrouwbaarheidscore raming", f"{_fmt_number(estimation.confidence_score, 0)}%"])
        summary_data.append(["Bandbreedte", f"{_fmt_currency(estimation.bandwidth_low)} — {_fmt_currency(estimation.bandwidth_high)}"])

    t_summary = Table(summary_data, colWidths=[80 * mm, 80 * mm])
    t_summary.setStyle(TableStyle(table_header_style))
    story.append(t_summary)

    if project.scope_description:
        story.append(Spacer(1, 4 * mm))
        story.append(Paragraph("Projectomschrijving", style_label))
        story.append(Paragraph(project.scope_description, style_body))

    # -------------------------------------------------------------------------
    # Hours per discipline
    # -------------------------------------------------------------------------
    story.append(Paragraph("Urenoverzicht per discipline", style_h2))

    if estimation:
        total_hours = sum(
            float(v or 0)
            for v in [
                estimation.hours_engineering,
                estimation.hours_pm,
                estimation.hours_om,
                estimation.hours_workprep,
                estimation.hours_execution,
            ]
        )
        hours_data = [
            ["Discipline / Rol", "Uren", "% van totaal"],
            [
                "Engineering",
                _fmt_number(estimation.hours_engineering),
                f"{float(estimation.hours_engineering or 0) / total_hours * 100:.1f}%" if total_hours else "—",
            ],
            [
                "Projectmanagement",
                _fmt_number(estimation.hours_pm),
                f"{float(estimation.hours_pm or 0) / total_hours * 100:.1f}%" if total_hours else "—",
            ],
            [
                "Omgevingsmanagement",
                _fmt_number(estimation.hours_om),
                f"{float(estimation.hours_om or 0) / total_hours * 100:.1f}%" if total_hours else "—",
            ],
            [
                "Werkvoorbereiding",
                _fmt_number(estimation.hours_workprep),
                f"{float(estimation.hours_workprep or 0) / total_hours * 100:.1f}%" if total_hours else "—",
            ],
            [
                "Uitvoering",
                _fmt_number(estimation.hours_execution),
                f"{float(estimation.hours_execution or 0) / total_hours * 100:.1f}%" if total_hours else "—",
            ],
            ["TOTAAL", _fmt_number(total_hours), "100%"],
        ]
        t_hours = Table(hours_data, colWidths=[100 * mm, 40 * mm, 40 * mm])
        style_hours = TableStyle(
            table_header_style
            + [
                ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor(f"#{NAVY_LIGHT}")),
                ("TEXTCOLOR", (0, -1), (-1, -1), white_c),
            ]
        )
        t_hours.setStyle(style_hours)
        story.append(t_hours)

        if estimation.reasoning:
            story.append(Spacer(1, 3 * mm))
            story.append(Paragraph("Onderbouwing", style_label))
            story.append(Paragraph(estimation.reasoning, style_body))
        story.append(Spacer(1, 3 * mm))
        story.append(Paragraph(f"Methodiek: {estimation.methodology or '—'}", style_small))
    else:
        story.append(Paragraph("Geen actuele raming beschikbaar.", style_body))

    # -------------------------------------------------------------------------
    # Risk matrix summary (top 10)
    # -------------------------------------------------------------------------
    story.append(Paragraph("Risicomatrix — Top 10 risico's", style_h2))

    if risks:
        risk_data = [["Nr.", "Beschrijving", "Cat.", "Kans", "Impact", "Score", "Status"]]
        for idx, risk in enumerate(risks[:10], 1):
            risk_data.append(
                [
                    str(idx),
                    Paragraph(risk.description[:100], style_small),
                    risk.category,
                    str(risk.probability),
                    str(risk.impact),
                    str(risk.score),
                    STATUS_NL.get(risk.status, risk.status),
                ]
            )

        t_risks = Table(
            risk_data,
            colWidths=[8 * mm, 75 * mm, 22 * mm, 12 * mm, 14 * mm, 12 * mm, 22 * mm],
        )
        risk_style = TableStyle(
            table_header_style
            + [
                ("WORDWRAP", (1, 1), (1, -1), True),
            ]
        )
        t_risks.setStyle(risk_style)
        story.append(t_risks)
    else:
        story.append(Paragraph("Geen risico's geregistreerd.", style_body))

    # -------------------------------------------------------------------------
    # Permits table
    # -------------------------------------------------------------------------
    story.append(Paragraph("Vergunningoverzicht", style_h2))

    if permits:
        permit_data = [["Type vergunning", "Bevoegd gezag", "Status", "Risico", "Verwachte goedkeuring"]]
        for permit in permits:
            permit_data.append(
                [
                    Paragraph(permit.permit_type, style_small),
                    permit.authority,
                    STATUS_NL.get(permit.status, permit.status),
                    (permit.risk_level or "").upper(),
                    _fmt_date(permit.expected_approval),
                ]
            )
        t_permits = Table(
            permit_data,
            colWidths=[65 * mm, 30 * mm, 28 * mm, 18 * mm, 35 * mm],
        )
        t_permits.setStyle(TableStyle(table_header_style))
        story.append(t_permits)
    else:
        story.append(Paragraph("Geen vergunningen geregistreerd.", style_body))

    story.append(Spacer(1, 6 * mm))
    story.append(
        Paragraph(
            f"<i>Dit rapport is automatisch gegenereerd op {report_date}. "
            "Controleer alle gegevens voor gebruik in formele procedures.</i>",
            style_small,
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer.read()
